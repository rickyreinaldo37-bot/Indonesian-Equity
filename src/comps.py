"""F2 — Comps engine.

Computes the sector comparables table from cached data + the manual metrics
layer and writes a formatted Excel workbook:

    output/comps_YYYY-MM-DD.xlsx
        Summary   — one row per bank + sector aggregates, styled
        <TICKER>  — one detail tab per bank (valuation, quarterly ratios,
                    annual financials)
        Notes     — sources, methodology, disclaimer

Run:  python comps.py
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src import config, fetch, flows, forwards, valuation

NAVY = "1F3864"
LIGHT = "F2F5F9"
RED = "C00000"
GRAY = "808080"

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
TOPLINE = Border(top=Side(style="medium", color=NAVY))


@dataclass
class CompsResult:
    table: pd.DataFrame        # index: ticker, columns: metrics
    aggregates: pd.DataFrame   # index: Median / Weighted mean
    as_of: dt.date
    manual_period: str         # latest manual quarter label (e.g. 2026Q1)
    fy_label: str              # fiscal year behind trailing metrics (e.g. FY25)
    sample: bool


# ---------------------------------------------------------------------------
# Metric computation
# ---------------------------------------------------------------------------

def _bank_row(ticker: str) -> dict:
    price, price_date = fetch.last_close(ticker)
    fund = fetch.fundamentals_frame(ticker)
    latest_fy = int(fund.index[-1])
    row_fy = fund.loc[latest_fy]
    prev_fy = fund.loc[latest_fy - 1] if (latest_fy - 1) in fund.index else row_fy

    shares = float(fetch.load_fundamentals(ticker)["shares_outstanding"])
    mcap_tn = price * shares / 1e12

    avg_equity = (row_fy["total_equity"] + prev_fy["total_equity"]) / 2
    avg_assets = (row_fy["total_assets"] + prev_fy["total_assets"]) / 2

    assumptions = valuation.load_assumptions(ticker)
    fwd_eps = assumptions.get("forward_eps")

    manual = fetch.latest_manual(ticker)
    px = fetch.load_prices(ticker)["Close"].dropna()
    yr = px[px.index >= px.index[-1] - pd.Timedelta(days=365)]
    flow = flows.bank_flow_summary(ticker)
    fwd = forwards.forward_view(ticker)

    return {
        "name": config.UNIVERSE[ticker]["short"],
        "price": price,
        "price_date": price_date.date(),
        "hi_52w": float(yr.max()),
        "lo_52w": float(yr.min()),
        "mcap_tn": mcap_tn,
        "pe_trailing": price / row_fy["eps"],
        "pe_forward": (price / fwd_eps) if fwd_eps else float("nan"),
        "pb": price / row_fy["bvps"],
        "roe_pct": 100 * row_fy["net_income"] / avg_equity,
        "roa_pct": 100 * row_fy["net_income"] / avg_assets,
        "div_yield_pct": 100 * row_fy["dps"] / price,
        "nim_pct": manual["nim_pct"],
        "casa_pct": manual["casa_pct"],
        "npl_pct": manual["npl_pct"],
        "car_pct": manual["car_pct"],
        "loan_growth_yoy_pct": manual["loan_growth_yoy_pct"],
        "cost_of_credit_pct": manual["cost_of_credit_pct"],
        "cost_income_pct": manual["cost_income_pct"],
        "flow_3m_tn": flow["flow_3m_tn"],
        "flow_12m_tn": flow["flow_12m_tn"],
        "flow_corr_1y": flow["corr_daily_1y"],
        "consensus_tp": fwd.target_mean if fwd.consensus_available
        else float("nan"),
        "consensus_upside_pct": fwd.upside_to_target_pct
        if fwd.consensus_available else float("nan"),
        "manual_period": manual["period"],
        "latest_fy": latest_fy,
        "eps": row_fy["eps"],
        "bvps": row_fy["bvps"],
        "dps": row_fy["dps"],
    }


NUMERIC_COLS = [
    "price", "mcap_tn", "pe_trailing", "pe_forward", "pb", "roe_pct",
    "roa_pct", "div_yield_pct", "nim_pct", "casa_pct", "npl_pct", "car_pct",
    "loan_growth_yoy_pct", "cost_of_credit_pct", "cost_income_pct",
    "flow_3m_tn", "flow_corr_1y", "consensus_tp", "consensus_upside_pct",
]


def build_comps() -> CompsResult:
    rows = {t: _bank_row(t) for t in config.TICKERS}
    table = pd.DataFrame(rows).T

    num = table[NUMERIC_COLS].astype(float)
    weights = num["mcap_tn"] / num["mcap_tn"].sum()
    aggregates = pd.DataFrame({
        "Median": num.median(),
        "Cap-weighted mean": (num.mul(weights, axis=0)).sum(),
    }).T
    # Weighted mean of price/52w levels is meaningless — blank them out.
    for col in ("price",):
        aggregates.loc["Cap-weighted mean", col] = float("nan")
        aggregates.loc["Median", col] = float("nan")
    # Cap-weighting a correlation coefficient is statistically meaningless,
    # and a cap-weighted mean of net flows has no economic reading either
    # (the meaningful sector aggregate — the total — goes in the footnote).
    aggregates.loc["Cap-weighted mean", "flow_corr_1y"] = float("nan")
    aggregates.loc["Cap-weighted mean", "flow_3m_tn"] = float("nan")
    # Target price is a per-share level, not comparable across banks — no
    # median or mean. Upside-to-target % is comparable; keep its median only.
    aggregates.loc[["Median", "Cap-weighted mean"], "consensus_tp"] = float("nan")
    aggregates.loc["Cap-weighted mean", "consensus_upside_pct"] = float("nan")

    # Premium / discount to sector median on the two valuation multiples.
    for col, out in (("pb", "pb_prem_disc_pct"),
                     ("pe_trailing", "pe_prem_disc_pct")):
        med = num[col].median()
        table[out] = 100 * (num[col] / med - 1)

    as_of = max(r["price_date"] for r in rows.values())
    fy = max(int(r["latest_fy"]) for r in rows.values())
    return CompsResult(
        table=table,
        aggregates=aggregates,
        as_of=as_of,
        manual_period=str(table["manual_period"].iloc[-1]),
        fy_label=f"FY{fy % 100}",
        sample=fetch.any_sample_data(),
    )


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

F_TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=9, color=GRAY)
F_WARN = Font(name="Calibri", size=9, bold=True, color=RED)
F_HEAD = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
F_BODY = Font(name="Calibri", size=10)
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_SECTION = Font(name="Calibri", size=10, bold=True, color=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_LIGHT = PatternFill("solid", fgColor=LIGHT)

FMT_PRICE = "#,##0"
FMT_TN = "#,##0.0"
FMT_TN_SIGNED = "+#,##0.0;-#,##0.0;0.0"
FMT_X = '0.0"x"'
FMT_PCT = '0.0"%"'
FMT_PCT2 = '0.00"%"'
FMT_CORR = "0.00"


def _title_block(ws, title: str, as_of: dt.date, sample: bool,
                 width: int) -> int:
    """Write title, as-of line and (if needed) sample banner. Returns the
    next free row."""
    ws.cell(row=1, column=1, value=title).font = F_TITLE
    src = ("SAMPLE / illustrative data — NOT for publication. "
           "Run refresh.py with live data before sharing."
           if sample else
           "Sources: Yahoo Finance (yfinance), company reports, author estimates")
    ws.cell(row=2, column=1,
            value=f"As of {as_of:%d %b %Y}  |  {src}  |  "
                  f"{config.WATERMARK}").font = (F_WARN if sample else F_SUB)
    for col in range(1, width + 1):
        ws.cell(row=3, column=col).border = TOPLINE
    return 4


SUMMARY_COLS: list[tuple[str, str, str]] = [
    # (df column, header, number format)
    ("name", "Company", ""),
    ("price", "Price\n(IDR)", FMT_PRICE),
    ("mcap_tn", "Mcap\n(IDR tn)", FMT_TN),
    ("pe_trailing", "P/E\n{fy}", FMT_X),
    ("pe_forward", "P/E\nFY+1E*", FMT_X),
    ("pb", "P/B", FMT_X),
    ("pb_prem_disc_pct", "P/B vs\nmedian", FMT_PCT),
    ("consensus_tp", "Cons.\nTP†", FMT_PRICE),
    ("consensus_upside_pct", "Upside\nto TP†", FMT_PCT),
    ("roe_pct", "ROE\n{fy}", FMT_PCT),
    ("roa_pct", "ROA\n{fy}", FMT_PCT2),
    ("div_yield_pct", "Div\nyield", FMT_PCT),
    ("nim_pct", "NIM\n{q}", FMT_PCT),
    ("casa_pct", "CASA\n{q}", FMT_PCT),
    ("npl_pct", "NPL\n{q}", FMT_PCT),
    ("car_pct", "CAR\n{q}", FMT_PCT),
    ("loan_growth_yoy_pct", "Loan gr.\nYoY", FMT_PCT),
    ("cost_of_credit_pct", "Cost of\ncredit", FMT_PCT2),
    ("cost_income_pct", "Cost/\nincome", FMT_PCT),
    ("flow_3m_tn", "Net frgn\n3M (tn)", FMT_TN_SIGNED),
    ("flow_corr_1y", "Flow/ret\ncorr 1Y", FMT_CORR),
]


def _write_summary(wb: Workbook, comps: CompsResult) -> None:
    ws = wb.active
    ws.title = "Summary"
    ncols = len(SUMMARY_COLS) + 1
    r0 = _title_block(ws, "IDX Banks — Sector Comps", comps.as_of,
                      comps.sample, ncols)

    # Header row
    ws.cell(row=r0, column=1, value="Ticker").font = F_HEAD
    ws.cell(row=r0, column=1).fill = FILL_HEAD
    for j, (_, header, _) in enumerate(SUMMARY_COLS, start=2):
        cell = ws.cell(row=r0, column=j,
                       value=header.format(fy=comps.fy_label,
                                           q=comps.manual_period))
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[r0].height = 26

    # Bank rows
    r = r0 + 1
    for ticker in config.TICKERS:
        row = comps.table.loc[ticker]
        ws.cell(row=r, column=1, value=ticker).font = F_BOLD
        for j, (col, _, fmt) in enumerate(SUMMARY_COLS, start=2):
            val = row[col]
            cell = ws.cell(row=r, column=j)
            if isinstance(val, float) and pd.isna(val):
                cell.value = "—"
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.value = float(val) if col != "name" else val
                if fmt:
                    cell.number_format = fmt
            cell.font = F_BODY
            cell.border = BORDER
        ws.cell(row=r, column=1).border = BORDER
        r += 1

    # Aggregate rows
    for label in ("Median", "Cap-weighted mean"):
        ws.cell(row=r, column=1, value=label).font = F_BOLD
        ws.cell(row=r, column=1).border = TOPLINE
        for j, (col, _, fmt) in enumerate(SUMMARY_COLS, start=2):
            cell = ws.cell(row=r, column=j)
            cell.border = TOPLINE
            if col in comps.aggregates.columns:
                val = comps.aggregates.loc[label, col]
                if pd.notna(val):
                    cell.value = float(val)
                    cell.number_format = fmt or "0.0"
                    cell.font = F_BOLD
        r += 1

    # Footnotes
    sector_3m = comps.table["flow_3m_tn"].astype(float).sum()
    ws.cell(row=r + 1, column=1,
            value="* Forward P/E uses the analyst's own FY+1 EPS estimate "
                  "from assumptions/<TICKER>.yaml (no consensus feed) — "
                  "clearly labeled, not consensus.").font = F_SUB
    ws.cell(row=r + 2, column=1,
            value=f"Net foreign flow: net foreign buy value, IDR tn "
                  f"(+ = inflow). Sector 3M total: {sector_3m:+.1f} tn. "
                  f"Corr = Pearson(daily net flow, daily return), trailing "
                  f"1Y — contemporaneous co-movement, not a forecast."
            ).font = F_SUB
    cons_srcs = ", ".join(sorted({fetch.consensus_source(t)
                                  for t in config.TICKERS}))
    ws.cell(row=r + 3, column=1,
            value=f"† Cons. TP = sell-side consensus mean target price; "
                  f"Upside to TP = vs last close. Source(s): {cons_srcs}. "
                  f"Consensus is never fabricated — n/a where no source "
                  f"exists; SAMPLE where illustrative."
            ).font = F_SUB

    # Conditional formatting: premium/discount to median (green=cheap,
    # red=expensive)
    pd_col = next(j for j, (c, _, _) in enumerate(SUMMARY_COLS, start=2)
                  if c == "pb_prem_disc_pct")
    letter = get_column_letter(pd_col)
    ws.conditional_formatting.add(
        f"{letter}{r0 + 1}:{letter}{r0 + len(config.TICKERS)}",
        ColorScaleRule(start_type="num", start_value=-40,
                       start_color="63BE7B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=40, end_color="F8696B"))
    # …and on 3M net foreign flow (red=outflow, green=inflow)
    fl_col = next(j for j, (c, _, _) in enumerate(SUMMARY_COLS, start=2)
                  if c == "flow_3m_tn")
    letter = get_column_letter(fl_col)
    ws.conditional_formatting.add(
        f"{letter}{r0 + 1}:{letter}{r0 + len(config.TICKERS)}",
        ColorScaleRule(start_type="num", start_value=-5, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=5, end_color="63BE7B"))
    # …and on upside to consensus target (red=downside, green=upside)
    up_col = next(j for j, (c, _, _) in enumerate(SUMMARY_COLS, start=2)
                  if c == "consensus_upside_pct")
    letter = get_column_letter(up_col)
    ws.conditional_formatting.add(
        f"{letter}{r0 + 1}:{letter}{r0 + len(config.TICKERS)}",
        ColorScaleRule(start_type="num", start_value=-20, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=20, end_color="63BE7B"))

    # Layout
    ws.freeze_panes = ws.cell(row=r0 + 1, column=3)  # freeze header + ticker+name
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 22
    for j in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 9.5
    ws.sheet_view.showGridLines = False


def _kv_rows(ws, r: int, pairs: list[tuple[str, object, str]],
             label_width: int = 1) -> int:
    """Write label/value rows; returns next free row."""
    for label, val, fmt in pairs:
        ws.cell(row=r, column=1, value=label).font = F_BODY
        ws.cell(row=r, column=1).fill = FILL_LIGHT
        cell = ws.cell(row=r, column=2)
        cell.value = val
        if fmt:
            cell.number_format = fmt
        cell.font = F_BODY
        r += 1
    return r


def _write_bank_tab(wb: Workbook, ticker: str, comps: CompsResult) -> None:
    ws = wb.create_sheet(ticker)
    row = comps.table.loc[ticker]
    info = config.UNIVERSE[ticker]
    r = _title_block(ws, f"{ticker} — {info['name']}", comps.as_of,
                     comps.sample, 12)

    # --- Valuation & market data block
    ws.cell(row=r, column=1, value="MARKET DATA & VALUATION").font = F_SECTION
    r += 1
    val = valuation.value_bank(ticker)
    pairs = [
        ("Last close (IDR)", float(row["price"]), FMT_PRICE),
        ("52-week high / low", f"{row['hi_52w']:,.0f} / {row['lo_52w']:,.0f}", ""),
        ("Market cap (IDR tn)", float(row["mcap_tn"]), FMT_TN),
        (f"EPS {comps.fy_label} (IDR)", float(row["eps"]), FMT_PRICE),
        (f"BVPS {comps.fy_label} (IDR)", float(row["bvps"]), FMT_PRICE),
        (f"DPS {comps.fy_label} (IDR)", float(row["dps"]), FMT_PRICE),
        (f"P/E {comps.fy_label}", float(row["pe_trailing"]), FMT_X),
        ("P/E FY+1E (own est.)", None if pd.isna(row["pe_forward"])
         else float(row["pe_forward"]), FMT_X),
        ("P/B", float(row["pb"]), FMT_X),
        ("Dividend yield", float(row["div_yield_pct"]), FMT_PCT),
        ("", "", ""),
        ("Residual income fair value (IDR)", round(val.ri_fair_value), FMT_PRICE),
        ("DDM cross-check (IDR)", round(val.ddm_fair_value), FMT_PRICE),
        ("Implied P/B at RI fair value", val.implied_pb, FMT_X),
        ("Upside / (downside)", val.upside_pct, FMT_PCT),
        ("Cost of equity (CAPM)", val.coe_pct, FMT_PCT2),
    ]
    r = _kv_rows(ws, r, pairs)
    r += 1

    # --- Forward estimates & consensus
    ws.cell(row=r, column=1,
            value="FORWARD ESTIMATES & CONSENSUS").font = F_SECTION
    r += 1
    fwd = forwards.forward_view(ticker)
    fwd_pairs = [
        ("Forward EPS FY+1E, own est. (IDR)",
         round(fwd.fwd_eps_own) if fwd.fwd_eps_own else "n/a",
         FMT_PRICE if fwd.fwd_eps_own else ""),
        ("Forward net income, own est. (IDR tn)",
         fwd.fwd_ni_own_tn if fwd.fwd_ni_own_tn else "n/a",
         FMT_TN if fwd.fwd_ni_own_tn else ""),
        ("EPS growth YoY, own est.",
         fwd.eps_growth_pct if fwd.eps_growth_pct is not None else "n/a",
         FMT_PCT if fwd.eps_growth_pct is not None else ""),
        ("Forward P/E, own est.",
         fwd.fwd_pe_own if fwd.fwd_pe_own else "n/a",
         FMT_X if fwd.fwd_pe_own else ""),
        ("RI-model year-1 EPS (cross-check)", round(fwd.model_year1_eps),
         FMT_PRICE),
        ("", "", ""),
    ]
    if fwd.consensus_available:
        tag = " [SAMPLE]" if fwd.consensus_is_sample else ""
        fwd_pairs += [
            (f"Consensus mean target (IDR){tag}", round(fwd.target_mean),
             FMT_PRICE),
            ("Upside to consensus target", fwd.upside_to_target_pct, FMT_PCT),
            ("Consensus analysts / rating",
             f"{fwd.num_analysts or 'n/a'} / {fwd.recommendation or 'n/a'}", ""),
            ("Consensus forward P/E",
             fwd.fwd_pe_consensus if fwd.fwd_pe_consensus else "n/a",
             FMT_X if fwd.fwd_pe_consensus else ""),
            ("Consensus source", str(fwd.consensus_source), ""),
        ]
    else:
        fwd_pairs.append(
            ("Consensus target", "n/a — no consensus source configured", ""))
    r = _kv_rows(ws, r, fwd_pairs)
    r += 1

    # --- Foreign flow monitor
    ws.cell(row=r, column=1,
            value="FOREIGN FLOW MONITOR (net foreign buy value)"
            ).font = F_SECTION
    r += 1
    fs = flows.bank_flow_summary(ticker)
    h = fs["horizons_tn"]
    flow_pairs = [
        (f"Net foreign flow {lbl} (IDR tn)", float(h[lbl]), FMT_TN_SIGNED)
        for lbl in ("1W", "1M", "3M", "6M", "YTD", "12M")
    ] + [
        ("Corr(daily flow, daily return), 1Y", fs["corr_daily_1y"], FMT_CORR),
        ("Corr(weekly flow, weekly return), 1Y", fs["corr_weekly_1y"],
         FMT_CORR),
        ("Flow data through", str(fs["last_date"]), ""),
    ]
    r = _kv_rows(ws, r, flow_pairs)
    r += 1

    # --- Quarterly manual metrics table
    ws.cell(row=r, column=1,
            value="KEY BANK METRICS BY QUARTER (manual layer — investor "
                  "presentations)").font = F_SECTION
    r += 1
    manual = fetch.load_manual_metrics(ticker).tail(8)
    metric_rows = [
        ("NIM", "nim_pct", FMT_PCT), ("CASA ratio", "casa_pct", FMT_PCT),
        ("NPL ratio", "npl_pct", FMT_PCT), ("CAR", "car_pct", FMT_PCT),
        ("Loan growth YoY", "loan_growth_yoy_pct", FMT_PCT),
        ("Cost of credit", "cost_of_credit_pct", FMT_PCT2),
        ("Cost/income", "cost_income_pct", FMT_PCT),
    ]
    ws.cell(row=r, column=1, value="Metric").font = F_HEAD
    ws.cell(row=r, column=1).fill = FILL_HEAD
    for j, (_, q) in enumerate(manual.iterrows(), start=2):
        c = ws.cell(row=r, column=j, value=q["period"])
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center")
    r += 1
    for label, col, fmt in metric_rows:
        ws.cell(row=r, column=1, value=label).font = F_BODY
        ws.cell(row=r, column=1).fill = FILL_LIGHT
        for j, (_, q) in enumerate(manual.iterrows(), start=2):
            c = ws.cell(row=r, column=j, value=float(q[col]))
            c.number_format = fmt
            c.font = F_BODY
            c.border = BORDER
        r += 1
    r += 1

    # --- Annual financials
    ws.cell(row=r, column=1,
            value="ANNUAL FINANCIALS (yfinance cache)").font = F_SECTION
    r += 1
    fund = fetch.fundamentals_frame(ticker)
    years = list(fund.index)[-6:]
    fin_rows = [
        ("Net income (IDR tn)", lambda y: fund.loc[y, "net_income"] / 1e12, FMT_TN),
        ("Total equity (IDR tn)", lambda y: fund.loc[y, "total_equity"] / 1e12, FMT_TN),
        ("Total assets (IDR tn)", lambda y: fund.loc[y, "total_assets"] / 1e12, FMT_TN),
        ("EPS (IDR)", lambda y: fund.loc[y, "eps"], FMT_PRICE),
        ("BVPS (IDR)", lambda y: fund.loc[y, "bvps"], FMT_PRICE),
        ("DPS (IDR)", lambda y: fund.loc[y, "dps"], FMT_PRICE),
        ("ROE (avg equity)", lambda y: _roe(fund, y), FMT_PCT),
        ("ROA (avg assets)", lambda y: _roa(fund, y), FMT_PCT2),
    ]
    ws.cell(row=r, column=1, value="FY").font = F_HEAD
    ws.cell(row=r, column=1).fill = FILL_HEAD
    for j, y in enumerate(years, start=2):
        c = ws.cell(row=r, column=j, value=f"FY{y % 100}")
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center")
    r += 1
    for label, getter, fmt in fin_rows:
        ws.cell(row=r, column=1, value=label).font = F_BODY
        ws.cell(row=r, column=1).fill = FILL_LIGHT
        for j, y in enumerate(years, start=2):
            v = getter(y)
            c = ws.cell(row=r, column=j)
            if v is not None:
                c.value = float(v)
                c.number_format = fmt
            c.font = F_BODY
            c.border = BORDER
        r += 1

    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 30
    for j in range(2, 12):
        ws.column_dimensions[get_column_letter(j)].width = 10.5
    ws.sheet_view.showGridLines = False


def _roe(fund: pd.DataFrame, year: int) -> float | None:
    prev = year - 1
    if prev not in fund.index:
        return None
    avg = (fund.loc[year, "total_equity"] + fund.loc[prev, "total_equity"]) / 2
    return 100 * fund.loc[year, "net_income"] / avg


def _roa(fund: pd.DataFrame, year: int) -> float | None:
    prev = year - 1
    if prev not in fund.index:
        return None
    avg = (fund.loc[year, "total_assets"] + fund.loc[prev, "total_assets"]) / 2
    return 100 * fund.loc[year, "net_income"] / avg


def _write_notes(wb: Workbook, comps: CompsResult) -> None:
    ws = wb.create_sheet("Notes")
    r = _title_block(ws, "Sources & methodology", comps.as_of, comps.sample, 8)
    lines = [
        "DATA SOURCES",
        "  Prices & fundamentals: Yahoo Finance via yfinance, cached locally "
        "(data/cache/). Source tag per file: "
        + ", ".join(f"{t}={fetch.price_source(t)}" for t in config.TICKERS),
        "  Bank-specific ratios (NIM, CASA, NPL, CAR, loan growth, cost of "
        "credit, cost/income): manual layer transcribed from quarterly "
        "investor presentations (data/manual/<TICKER>.csv) — source of truth.",
        "  Net foreign flow: "
        + ", ".join(f"{t}={fetch.flow_source(t)}" for t in config.TICKERS)
        + " (manual CSV export > IDX daily trading summary > cache).",
        "  Macro: data/manual/macro.csv (BI rate, USD/IDR, Indonesia 10Y).",
        "",
        "METHODOLOGY",
        f"  Trailing multiples use last close vs latest reported fiscal year "
        f"({comps.fy_label}) per-share values.",
        "  Forward P/E uses the analyst's own FY+1 EPS estimate from "
        "assumptions/<TICKER>.yaml — labeled 'own est.', no consensus feed.",
        "  ROE/ROA use average of opening and closing equity/assets.",
        "  Sector aggregates: median and market-cap-weighted mean.",
        "  Fair values: residual income model (primary) with DDM cross-check; "
        "see per-bank tabs and assumptions/<TICKER>.yaml.",
        "",
        "DISCLAIMER",
        f"  {config.DISCLAIMER}",
    ]
    if comps.sample:
        lines = ["*** SAMPLE DATA MODE — every figure in this workbook is "
                 "illustrative. Do not publish. ***", ""] + lines
    for line in lines:
        cell = ws.cell(row=r, column=1, value=line)
        cell.font = F_WARN if line.startswith("***") else (
            F_SECTION if line.isupper() and line else F_BODY)
        r += 1
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False


def write_excel(comps: CompsResult | None = None) -> str:
    comps = comps or build_comps()
    wb = Workbook()
    _write_summary(wb, comps)
    for ticker in config.TICKERS:
        _write_bank_tab(wb, ticker, comps)
    _write_notes(wb, comps)
    out = config.OUTPUT_DIR / f"comps_{comps.as_of:%Y-%m-%d}.xlsx"
    wb.save(out)
    return str(out)


def main() -> int:
    comps = build_comps()
    cols = ["price", "mcap_tn", "pe_trailing", "pe_forward", "pb", "roe_pct",
            "div_yield_pct", "nim_pct", "npl_pct"]
    print(f"IDX Banks comps — as of {comps.as_of}")
    with pd.option_context("display.float_format", "{:,.1f}".format):
        print(comps.table[cols].astype(float).to_string())
        print("\nSector aggregates:")
        print(comps.aggregates[cols].astype(float).to_string())
    path = write_excel(comps)
    print(f"\nWorkbook written: {path}")
    if comps.sample:
        print("NOTE: built from SAMPLE data — see Notes tab banner.")
    return 0


if __name__ == "__main__":
    main()
