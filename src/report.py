"""F5 — Report skeleton generator.

Produces an initiation-of-coverage skeleton with the standard sell-side
structure. Live numbers, valuation tables, and charts are injected
automatically; the qualitative sections (thesis, overview, risks) are
placeholders for the analyst to write.

    python make_report.py BBRI              → docx
    python make_report.py BBRI --format md  → markdown

Everything numeric flows from the caches + assumptions/<TICKER>.yaml, so
changing an assumption and re-running rebuilds the whole report consistently.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src import charts, config, fetch, flows, valuation
from src.comps import build_comps

NAVY = RGBColor(0x1F, 0x38, 0x64)
RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

PLACEHOLDER_THESIS = [
    "Thesis point 1 — the core reason to own (or avoid) the stock. "
    "State it in one sentence, then support it with two or three "
    "paragraphs of evidence.",
    "Thesis point 2 — what the market is missing or mispricing.",
    "Thesis point 3 — the catalyst path and timing.",
]
PLACEHOLDER_RISKS = [
    "Macro: BI policy path, IDR volatility, sovereign yield moves.",
    "Asset quality: NPL formation and cost of credit vs assumptions.",
    "Competition/funding: deposit pricing pressure on NIM.",
    "Regulatory/political: SoE dividend policy, directed lending.",
]


def _idx_tick(price: float) -> int:
    if price < 200:
        return 1
    if price < 500:
        return 2
    if price < 2000:
        return 5
    if price < 5000:
        return 10
    return 25


def round_to_tick(price: float) -> float:
    tick = _idx_tick(price)
    return round(price / tick) * tick


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------

def _base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(1.8)


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _para(doc, text="", size=10, bold=False, color=None, align=None,
          space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _banner(doc, text: str, fill: str, color=WHITE, size=11) -> None:
    """Full-width single-cell shaded banner row."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _shade(cell, fill)
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _df_table(doc, df: pd.DataFrame, index_header: str = "",
              highlight: tuple[int, int] | None = None) -> None:
    """Render a DataFrame (values already formatted as strings) as a compact
    styled table. highlight = (row, col) in df coordinates."""
    tbl = doc.add_table(rows=len(df) + 1, cols=len(df.columns) + 1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # header row
    hdr = tbl.rows[0]
    for j, text in enumerate([index_header] + [str(c) for c in df.columns]):
        cell = hdr.cells[j]
        _shade(cell, "1F3864")
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = WHITE
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # body
    for i, (idx, row) in enumerate(df.iterrows(), start=1):
        cells = tbl.rows[i].cells
        p = cells[0].paragraphs[0]
        run = p.add_run(str(idx))
        run.font.size = Pt(8.5)
        run.font.bold = True
        for j, val in enumerate(row, start=1):
            p = cells[j].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(8.5)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if highlight and (i - 1, j - 1) == highlight:
                _shade(cells[j], "FFF2CC")
                run.font.bold = True


def _footer(doc) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{config.WATERMARK}  |  {dt.date.today():%d %b %Y}  |  "
                    f"Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def f_idr(v) -> str:
    return f"{v:,.0f}"


def f_x(v) -> str:
    return f"{v:.1f}x"


def f_pct(v, dp=1) -> str:
    return f"{v:.{dp}f}%"


# ---------------------------------------------------------------------------
# Report assembly
# ---------------------------------------------------------------------------

def _gather(ticker: str) -> dict:
    comps = build_comps()
    val = valuation.value_bank(ticker)
    chart_paths = charts.generate_all(comps.table)
    return {"comps": comps, "val": val, "charts": chart_paths,
            "row": comps.table.loc[ticker],
            "manual": fetch.load_manual_metrics(ticker),
            "fund": fetch.fundamentals_frame(ticker),
            "sample": comps.sample}


def _cover(doc, ticker: str, ctx: dict) -> None:
    info = config.UNIVERSE[ticker]
    val, row = ctx["val"], ctx["row"]
    a = val.assumptions
    rating = a.get("rating") or "[SET RATING — assumptions yaml]"
    pt = round_to_tick(val.ri_fair_value)

    _banner(doc, f"EQUITY RESEARCH  |  INITIATION OF COVERAGE  |  "
                 f"{dt.date.today():%d %B %Y}", "1F3864")
    if ctx["sample"]:
        _banner(doc, "DRAFT BUILT ON SAMPLE DATA — ILLUSTRATIVE ONLY, "
                     "NOT FOR PUBLICATION", "C00000")
    _para(doc)
    _para(doc, info["name"], size=20, bold=True, color=NAVY, space_after=0)
    _para(doc, f"{info['yf']}  |  Indonesia — Banks", size=11, color=GRAY)

    stats = pd.DataFrame({
        "": [
            ("Rating", rating),
            ("Price target (RI fair value)", f"IDR {f_idr(pt)}"),
            (f"Last close ({val.price_date})", f"IDR {f_idr(val.last_close)}"),
            ("Upside / (downside)", f_pct(val.upside_pct, 1)),
            ("Market cap", f"IDR {row['mcap_tn']:.0f} tn"),
            ("52-week range", f"{f_idr(row['lo_52w'])} – {f_idr(row['hi_52w'])}"),
        ],
        " ": [
            (f"P/E ({ctx['comps'].fy_label})", f_x(row["pe_trailing"])),
            ("P/E (FY+1E, own est.)",
             "n/a" if pd.isna(row["pe_forward"]) else f_x(row["pe_forward"])),
            ("P/B", f_x(row["pb"])),
            ("Dividend yield", f_pct(row["div_yield_pct"])),
            ("ROE / ROA", f"{f_pct(row['roe_pct'])} / "
                          f"{f_pct(row['roa_pct'], 2)}"),
            ("Cost of equity (CAPM)", f_pct(val.coe_pct, 2)),
        ],
    })
    tbl = doc.add_table(rows=6, cols=4)
    tbl.style = "Table Grid"
    for i in range(6):
        for c, pair_col in ((0, ""), (2, " ")):
            label, value = stats[pair_col].iloc[i]
            lab_cell, val_cell = tbl.rows[i].cells[c], tbl.rows[i].cells[c + 1]
            _shade(lab_cell, "F2F5F9")
            run = lab_cell.paragraphs[0].add_run(label)
            run.font.size = Pt(9)
            run = val_cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            run.font.bold = True
    _para(doc)
    _para(doc, f"Analyst: {config.ANALYST_NAME}  |  {config.SITE_NAME}",
          size=9, color=GRAY)


def _valuation_section(doc, ctx: dict) -> None:
    val = ctx["val"]
    a = val.assumptions
    ri = a["residual_income"]
    coe_a = a["cost_of_equity"]

    doc.add_heading("Valuation", level=1)
    if coe_a.get("coe_override_pct") is not None:
        coe_text = (f"Cost of equity is set explicitly at "
                    f"{f_pct(val.coe_pct, 2)} (override).")
    else:
        coe_text = (f"Cost of equity of {f_pct(val.coe_pct, 2)} is CAPM-derived: "
                    f"risk-free {f_pct(coe_a['risk_free_pct'], 1)} (Indonesia 10Y) "
                    f"+ beta {coe_a['beta']:.2f} × equity risk premium "
                    f"{f_pct(coe_a['equity_risk_premium_pct'], 1)}.")
    _para(doc,
          f"Primary methodology is a residual income model — appropriate for "
          f"banks, where value accrues on book equity. {coe_text} ROE fades "
          f"linearly from {f_pct(ri['roe_start_pct'])} to "
          f"{f_pct(ri['roe_terminal_pct'])} over {ri['forecast_years']} years "
          f"on a {f_pct(ri['payout_ratio_pct'], 0)} payout; beyond the horizon, "
          f"residual income grows at {f_pct(ri['terminal_growth_pct'])} with the "
          f"sustainable-growth payout of {f_pct(val.payout_terminal_pct)}. "
          f"All inputs live in assumptions/{val.ticker}.yaml.")

    fv_line = (f"Residual income fair value: IDR {f_idr(val.ri_fair_value)} "
               f"per share (implied P/B {val.implied_pb:.2f}x, justified "
               f"terminal P/B {val.justified_terminal_pb:.2f}x). DDM "
               f"cross-check: IDR {f_idr(val.ddm_fair_value)}. "
               f"Upside/(downside) vs last close: {val.upside_pct:+.1f}%.")
    _para(doc, fv_line, bold=True)

    doc.add_heading("Residual income forecast path", level=2)
    path = val.ri_path.copy()
    show = pd.DataFrame({
        "ROE": path["roe_pct"].map(lambda v: f_pct(v)),
        "BVPS open": path["bvps_open"].map(f_idr),
        "EPS": path["eps"].map(f_idr),
        "DPS": path["dps"].map(f_idr),
        "Residual income": path["ri"].map(f_idr),
        "PV of RI": path["pv_ri"].map(f_idr),
    }, index=[f"Year {int(y)}" for y in path["year"]])
    _df_table(doc, show)
    _para(doc, f"PV of explicit residual income "
               f"{f_idr(val.ri_path['pv_ri'].sum())} + PV of terminal value "
               f"{f_idr(val.terminal['pv_terminal'])} + opening BVPS "
               f"{f_idr(val.bvps_0)} = fair value "
               f"{f_idr(val.ri_fair_value)}.", size=8.5, color=GRAY)

    doc.add_heading("Sensitivity — fair value: COE × terminal ROE", level=2)
    sens = val.sensitivity.copy()
    fmt = sens.map(lambda v: f_idr(v) if pd.notna(v) else "n.m.")
    center = (sens.shape[0] // 2, sens.shape[1] // 2)
    _df_table(doc, fmt, index_header="COE \\ terminal ROE", highlight=center)
    _para(doc, "Shaded cell = base case.", size=8.5, color=GRAY)


def _flow_section(doc, ticker: str, ctx: dict) -> None:
    doc.add_heading("Foreign flow monitor", level=1)
    fs = flows.bank_flow_summary(ticker)
    h = fs["horizons_tn"]
    _para(doc,
          f"IDX blue chips trade with foreign money, and {ticker} is no "
          f"exception: over the trailing year the correlation between daily "
          f"net foreign flow and daily returns is {fs['corr_daily_1y']:.2f} "
          f"({fs['corr_weekly_1y']:.2f} on weekly data). Net foreign flow is "
          f"{h['3M']:+.1f} IDR tn over 3M and {h['12M']:+.1f} IDR tn over "
          f"12M (data through {fs['last_date']}). Correlation is "
          f"contemporaneous co-movement, not a forecasting claim.")
    tbl = pd.DataFrame(
        {lbl: [f"{h[lbl]:+.1f}"] for lbl in
         ("1W", "1M", "3M", "6M", "YTD", "12M")},
        index=["Net foreign flow (IDR tn)"])
    _df_table(doc, tbl, index_header="")
    _para(doc, "[TO WRITE — interpret positioning: who has been "
               "accumulating/distributing, and what would turn the flow.]",
          color=RED)


def _exhibits_section(doc, ticker: str, ctx: dict) -> None:
    doc.add_heading("Key exhibits", level=1)
    order = [
        (f"pb_band_{ticker}", None),
        (f"foreign_flow_{ticker}", None),
        ("foreign_flow_sector", None),
        ("price_rebased", None),
        ("roe_vs_pb", None),
        ("nim_trend", None),
        ("loan_growth_trend", None),
    ]
    for name, _ in order:
        path = ctx["charts"].get(name)
        if path and Path(path).exists():
            doc.add_picture(str(path), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _financials_section(doc, ticker: str, ctx: dict) -> None:
    doc.add_heading("Financials appendix", level=1)
    fund = ctx["fund"].tail(6)
    years = [f"FY{int(y) % 100}" for y in fund.index]

    def roe_row():
        vals = []
        for y in fund.index:
            prev = y - 1
            if prev in ctx["fund"].index:
                avg = (fund.loc[y, "total_equity"]
                       + ctx["fund"].loc[prev, "total_equity"]) / 2
                vals.append(f_pct(100 * fund.loc[y, "net_income"] / avg))
            else:
                vals.append("—")
        return vals

    annual = pd.DataFrame({
        "Net income (IDR tn)": (fund["net_income"] / 1e12).map("{:.1f}".format),
        "Total equity (IDR tn)": (fund["total_equity"] / 1e12).map("{:.0f}".format),
        "Total assets (IDR tn)": (fund["total_assets"] / 1e12).map("{:.0f}".format),
        "EPS (IDR)": fund["eps"].map(f_idr),
        "BVPS (IDR)": fund["bvps"].map(f_idr),
        "DPS (IDR)": fund["dps"].map(f_idr),
        "ROE": roe_row(),
    }, index=years).T
    _df_table(doc, annual, index_header="Annual")
    _para(doc)

    doc.add_heading("Quarterly bank metrics (manual layer)", level=2)
    manual = ctx["manual"].tail(8)
    quarterly = pd.DataFrame({
        "NIM": manual["nim_pct"].map(lambda v: f_pct(v)),
        "CASA": manual["casa_pct"].map(lambda v: f_pct(v)),
        "NPL": manual["npl_pct"].map(lambda v: f_pct(v)),
        "CAR": manual["car_pct"].map(lambda v: f_pct(v)),
        "Loan growth YoY": manual["loan_growth_yoy_pct"].map(lambda v: f_pct(v)),
        "Cost of credit": manual["cost_of_credit_pct"].map(lambda v: f_pct(v, 2)),
        "Cost/income": manual["cost_income_pct"].map(lambda v: f_pct(v)),
    }, index=list(manual["period"])).T
    _df_table(doc, quarterly, index_header="Metric")


def build_docx(ticker: str) -> str:
    ctx = _gather(ticker)
    doc = Document()
    _base_styles(doc)
    _footer(doc)

    _cover(doc, ticker, ctx)

    doc.add_heading("Investment thesis", level=1)
    _para(doc, "[TO WRITE — the system deliberately does not write this. "
               "Suggested structure below.]", color=RED)
    for i, stub in enumerate(PLACEHOLDER_THESIS, 1):
        _para(doc, f"{i}. {stub}", color=GRAY)

    doc.add_heading("Company overview", level=1)
    _para(doc, "[TO WRITE — franchise, segment mix, funding profile, "
               "management, ownership.]", color=RED)

    _valuation_section(doc, ctx)
    _flow_section(doc, ticker, ctx)
    _exhibits_section(doc, ticker, ctx)
    _financials_section(doc, ticker, ctx)

    doc.add_heading("Risks", level=1)
    _para(doc, "[TO WRITE — tailor to the name; generic sector list below.]",
          color=RED)
    for stub in PLACEHOLDER_RISKS:
        _para(doc, f"•  {stub}", color=GRAY)

    doc.add_heading("Disclosures", level=1)
    _para(doc, config.DISCLAIMER, size=8.5, color=GRAY)
    _para(doc, "Data: Yahoo Finance via yfinance (prices, financial "
               "statements); company quarterly investor presentations "
               "(bank-specific ratios, transcribed by the analyst); "
               "IDX daily trading summary / broker exports (net foreign "
               "flow); Bank Indonesia / market data (macro). Forward EPS "
               "estimates are the analyst's own — no consensus feed is used.",
          size=8.5, color=GRAY)
    if ctx["sample"]:
        _para(doc, "THIS DRAFT WAS BUILT ON SAMPLE PLACEHOLDER DATA — "
                   "figures are illustrative and must not be published.",
              size=9, bold=True, color=RED)

    out = (config.REPORT_DIR
           / f"{ticker}_initiation_skeleton_{dt.date.today():%Y-%m-%d}.docx")
    doc.save(out)
    return str(out)


# ---------------------------------------------------------------------------
# Markdown variant
# ---------------------------------------------------------------------------

def build_markdown(ticker: str) -> str:
    ctx = _gather(ticker)
    info = config.UNIVERSE[ticker]
    val, row = ctx["val"], ctx["row"]
    a = val.assumptions
    ri = a["residual_income"]
    pt = round_to_tick(val.ri_fair_value)
    rating = a.get("rating") or "[SET RATING]"
    lines: list[str] = []
    add = lines.append

    add(f"# {info['name']} ({info['yf']}) — Initiation of Coverage")
    add(f"*{dt.date.today():%d %B %Y} · {config.ANALYST_NAME} · "
        f"{config.SITE_NAME}*")
    if ctx["sample"]:
        add("\n> **DRAFT BUILT ON SAMPLE DATA — ILLUSTRATIVE ONLY, NOT FOR "
            "PUBLICATION**")
    add("\n| | | | |\n|---|---|---|---|")
    add(f"| **Rating** | {rating} | **P/E ({ctx['comps'].fy_label})** | "
        f"{f_x(row['pe_trailing'])} |")
    add(f"| **Price target (RI)** | IDR {f_idr(pt)} | **P/B** | "
        f"{f_x(row['pb'])} |")
    add(f"| **Last close** | IDR {f_idr(val.last_close)} | **Div yield** | "
        f"{f_pct(row['div_yield_pct'])} |")
    add(f"| **Upside/(downside)** | {val.upside_pct:+.1f}% | **ROE** | "
        f"{f_pct(row['roe_pct'])} |")
    add(f"| **Market cap** | IDR {row['mcap_tn']:.0f} tn | **COE (CAPM)** | "
        f"{f_pct(val.coe_pct, 2)} |")

    add("\n## Investment thesis\n")
    add("*[TO WRITE — the system deliberately does not write this.]*\n")
    for i, stub in enumerate(PLACEHOLDER_THESIS, 1):
        add(f"{i}. {stub}")

    add("\n## Company overview\n\n*[TO WRITE]*")

    add("\n## Valuation\n")
    add(f"Residual income model: COE {f_pct(val.coe_pct, 2)}, ROE fade "
        f"{f_pct(ri['roe_start_pct'])} → {f_pct(ri['roe_terminal_pct'])} over "
        f"{ri['forecast_years']}y, payout {f_pct(ri['payout_ratio_pct'], 0)}, "
        f"terminal growth {f_pct(ri['terminal_growth_pct'])}.\n")
    add(f"**Fair value IDR {f_idr(val.ri_fair_value)}** (implied P/B "
        f"{val.implied_pb:.2f}x) · DDM cross-check IDR "
        f"{f_idr(val.ddm_fair_value)} · upside {val.upside_pct:+.1f}%\n")
    path = val.ri_path
    add("| Year | ROE | BVPS open | EPS | DPS | RI | PV(RI) |")
    add("|---|---|---|---|---|---|---|")
    for _, r in path.iterrows():
        add(f"| {int(r['year'])} | {f_pct(r['roe_pct'])} | "
            f"{f_idr(r['bvps_open'])} | {f_idr(r['eps'])} | "
            f"{f_idr(r['dps'])} | {f_idr(r['ri'])} | {f_idr(r['pv_ri'])} |")
    add("\n**Sensitivity (fair value, COE × terminal ROE)**\n")
    sens = val.sensitivity
    add("| COE \\ ROE | " + " | ".join(sens.columns) + " |")
    add("|" + "---|" * (len(sens.columns) + 1))
    for idx, r in sens.iterrows():
        add(f"| {idx} | " + " | ".join(
            f_idr(v) if pd.notna(v) else "n.m." for v in r) + " |")

    fs = flows.bank_flow_summary(ticker)
    h = fs["horizons_tn"]
    add("\n## Foreign flow monitor\n")
    add(f"Trailing-1Y corr(daily net foreign flow, daily return): "
        f"**{fs['corr_daily_1y']:.2f}** (weekly {fs['corr_weekly_1y']:.2f}). "
        f"Data through {fs['last_date']}. Contemporaneous co-movement, not "
        f"a forecast.\n")
    add("| 1W | 1M | 3M | 6M | YTD | 12M |")
    add("|---|---|---|---|---|---|")
    add("| " + " | ".join(f"{h[lbl]:+.1f}" for lbl in
                          ("1W", "1M", "3M", "6M", "YTD", "12M"))
        + " |")
    add("\n*Net foreign buy value, IDR tn. [TO WRITE — interpret "
        "positioning.]*")

    add("\n## Key exhibits\n")
    rel = {
        f"pb_band_{ticker}": f"../charts/{ticker}/pb_band.png",
        f"foreign_flow_{ticker}": f"../charts/{ticker}/foreign_flow.png",
        "foreign_flow_sector": "../charts/sector/foreign_flow_cumulative.png",
        "price_rebased": "../charts/sector/price_rebased.png",
        "roe_vs_pb": "../charts/sector/roe_vs_pb.png",
        "nim_trend": "../charts/sector/nim_trend.png",
        "loan_growth_trend": "../charts/sector/loan_growth_trend.png",
    }
    for name, relpath in rel.items():
        add(f"![{name}]({relpath})")

    add("\n## Risks\n\n*[TO WRITE — generic sector list:]*\n")
    for stub in PLACEHOLDER_RISKS:
        add(f"- {stub}")

    add(f"\n---\n*{config.DISCLAIMER}*")
    if ctx["sample"]:
        add("\n**THIS DRAFT WAS BUILT ON SAMPLE PLACEHOLDER DATA — figures "
            "are illustrative and must not be published.**")

    out = (config.REPORT_DIR
           / f"{ticker}_initiation_skeleton_{dt.date.today():%Y-%m-%d}.md")
    Path(out).write_text("\n".join(lines))
    return str(out)


def build_report(ticker: str, fmt: str = "docx") -> str:
    if fmt == "md":
        return build_markdown(ticker)
    return build_docx(ticker)
